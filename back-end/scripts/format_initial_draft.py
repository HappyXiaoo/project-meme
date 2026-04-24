from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[2]
RES_DIR = ROOT / "res"


def pick_draft_file() -> Path:
    candidates = sorted(RES_DIR.glob("*.docx"), key=lambda p: p.stat().st_mtime, reverse=True)
    for path in candidates:
        if path.name.endswith(".docx") and path.stat().st_size > 50000:
            return path
    raise FileNotFoundError("No draft docx found in res/")


def set_run_font(run, font_name: str, size_pt: float, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold


def set_paragraph_runs(paragraph, font_name: str, size_pt: float, bold: bool = False) -> None:
    if not paragraph.runs:
        run = paragraph.add_run("")
        set_run_font(run, font_name, size_pt, bold)
        return
    for run in paragraph.runs:
        set_run_font(run, font_name, size_pt, bold)


def apply_page_setup(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(3)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.5)


def format_doc(doc: Document) -> None:
    apply_page_setup(doc)

    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        fmt = paragraph.paragraph_format
        fmt.line_spacing = 1.5
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)

        if not text:
            continue

        if idx == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_runs(paragraph, "宋体", 15, False)
            continue

        if text == "摘要":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_runs(paragraph, "黑体", 14, True)
            continue

        if text == "关键词：":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_runs(paragraph, "宋体", 14, True)
            continue

        if text == "Abstract":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_runs(paragraph, "Times New Roman", 14, True)
            continue

        if text == "Keywords:":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_runs(paragraph, "Times New Roman", 12, True)
            continue

        if text.startswith("Research and Design of"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_runs(paragraph, "Times New Roman", 15, False)
            continue

        if text.startswith("Author:") or text.startswith("Supervisor:"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            set_paragraph_runs(paragraph, "Times New Roman", 12, False)
            continue

        if text == "目    录":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_runs(paragraph, "黑体", 16, True)
            continue

        if text.startswith("第") and "章" in text and len(text) < 30:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_runs(paragraph, "宋体", 16, True)
            continue

        if re.match(r"^\d+\.\d+", text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_runs(paragraph, "宋体", 14, True)
            continue

        if text == "参考文献" or text == "致    谢":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_runs(paragraph, "宋体", 16, True)
            continue

        if text.startswith("表4-1"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_paragraph_runs(paragraph, "黑体", 10, False)
            continue

        if text.startswith("[") and "]" in text:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.first_line_indent = Pt(0)
            set_paragraph_runs(paragraph, "宋体", 10, False)
            continue

        if all(ch in ". 0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz第章节目录参考文献致谢注：上下左右" for ch in text) and len(text) > 10:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_runs(paragraph, "宋体", 12, False)
            continue

        # English abstract/body
        if re.match(r"^[A-Za-z].*", text):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.first_line_indent = Pt(0)
            set_paragraph_runs(paragraph, "Times New Roman", 12, False)
            continue

        # Default Chinese body
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt.first_line_indent = Pt(24)
        set_paragraph_runs(paragraph, "宋体", 12, False)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.line_spacing = 1.2
                    for run in paragraph.runs:
                        set_run_font(run, "宋体", 10, False)


def main() -> None:
    draft = pick_draft_file()
    doc = Document(str(draft))
    format_doc(doc)
    doc.save(str(draft))
    print(draft)


if __name__ == "__main__":
    main()
