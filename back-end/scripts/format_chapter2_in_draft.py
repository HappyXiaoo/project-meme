import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
RES_DIR = ROOT / "res"


def set_run_font(run, font_name: str, size_pt: float, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold


def apply_style(paragraph, *, font_name: str, size_pt: float, bold: bool, align, first_line_indent_pt: float = 0) -> None:
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(first_line_indent_pt)
    if not paragraph.runs:
        run = paragraph.add_run("")
        set_run_font(run, font_name, size_pt, bold)
    for run in paragraph.runs:
        set_run_font(run, font_name, size_pt, bold)


def find_draft() -> Path:
    files = sorted(RES_DIR.glob("*.docx"), key=lambda p: p.stat().st_size, reverse=True)
    if len(files) < 2:
        raise RuntimeError("draft docx not found")
    # current draft is the middle-sized docx in res
    return files[1]


def find_chapter2_range(doc: Document) -> tuple[int, int]:
    chapter2_heading = None
    chapter3_heading = None

    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if re.match(r"^2\.1\s", text):
            # chapter heading is previous non-empty paragraph
            j = i - 1
            while j >= 0 and not doc.paragraphs[j].text.strip():
                j -= 1
            chapter2_heading = j
            break

    if chapter2_heading is None:
        raise RuntimeError("chapter 2 start not found")

    for i in range(chapter2_heading + 1, len(doc.paragraphs)):
        text = doc.paragraphs[i].text.strip()
        if text.startswith("第3章"):
            chapter3_heading = i
            break

    if chapter3_heading is None:
        chapter3_heading = len(doc.paragraphs)

    return chapter2_heading, chapter3_heading


def main() -> None:
    draft = find_draft()
    doc = Document(str(draft))
    start, end = find_chapter2_range(doc)

    for i in range(start, end):
        p = doc.paragraphs[i]
        text = p.text.strip()
        if not text:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            continue

        if text.startswith("第2章"):
            apply_style(
                p,
                font_name="SimSun",
                size_pt=16,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
        elif re.match(r"^2\.\d+\.\d+\s", text):
            apply_style(
                p,
                font_name="SimSun",
                size_pt=12,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.LEFT,
            )
        elif re.match(r"^2\.\d+\s", text):
            apply_style(
                p,
                font_name="SimSun",
                size_pt=14,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.LEFT,
            )
        elif re.match(r"^[（(]\d+[)）]", text) or re.match(r"^\d+\.", text):
            apply_style(
                p,
                font_name="SimSun",
                size_pt=12,
                bold=False,
                align=WD_ALIGN_PARAGRAPH.LEFT,
            )
        else:
            apply_style(
                p,
                font_name="SimSun",
                size_pt=12,
                bold=False,
                align=WD_ALIGN_PARAGRAPH.LEFT,
                first_line_indent_pt=24,
            )

    doc.save(str(draft))
    print(draft)


if __name__ == "__main__":
    main()
