from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
RES_DIR = ROOT / "res"


def set_run_font(run, font_name: str, size_pt: float, bold: bool = False) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold


def format_paragraph(
    paragraph,
    *,
    font_name="SimSun",
    size_pt=12,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    first_line_indent_pt=0,
):
    paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Pt(first_line_indent_pt)
    if not paragraph.runs:
        paragraph.add_run("")
    for run in paragraph.runs:
        set_run_font(run, font_name, size_pt, bold)


def find_draft() -> Path:
    for path in RES_DIR.glob("*.docx"):
        if path.name.startswith("~$"):
            continue
        doc = Document(str(path))
        texts = [p.text.strip() for p in doc.paragraphs[:120]]
        if "2.1 系统需求分析" in texts:
            return path
    raise RuntimeError("未找到包含第2章内容的初稿文件")


def build_chapter3_items():
    return [
        (
            "3.1 数据集与预处理",
            "本研究使用的数据集以中文文本及其对应标签为基础，其中标签 0 表示正常文本，标签 1 表示疑似有害文本。"
            "在模型训练前，首先对原始数据进行编码读取与字段检查，保留标签列和文本列两个核心字段；随后对空文本、非法标签和重复文本进行清洗，"
            "以保证训练样本具有基本质量。对于传统机器学习模型，本文将清洗后的数据直接用于训练和测试；对于 BERT 模型，则进一步划分训练集、验证集和测试集，"
            "并在训练脚本中完成分批读取和 tokenizer 编码。除正式训练数据外，本文还构造了一份人工测试集，用于分析模型在模因式表达、隐含攻击表达和反对语境表达等样例上的具体表现。"
        ),
        (
            "3.2 基线模型实现",
            "基线模型采用字符级 TF-IDF 与 Logistic Regression 的组合方式实现。之所以选择该方案，主要是因为其实现简单、训练效率高、结果稳定，"
            "适合作为整篇论文实验比较的初始参照模型。具体而言，系统先通过 TF-IDF 将中文文本转换为字符级 n-gram 特征，再利用 Logistic Regression 完成二分类判断。"
            "该模型对显式攻击文本具有较好的识别能力，也能够作为后续人工特征模型和 BERT 模型的性能基线。为了便于系统调用和实验复现，"
            "本文将训练后的模型保存为本地文件，并在后端服务启动时优先加载。"
        ),
        (
            "3.3 模因式表达特征模型实现",
            "为了增强系统对文本模因场景的适应能力，本文在基线模型基础上增加了一组人工设计的模因式表达特征，并构建了模因式表达特征增强模型。"
            "该模型保留了字符级 TF-IDF 作为基础表示，同时加入模板化表达、群体指向、情绪强化、网络化短语、讽刺表达、风险词与群体词共现以及语境消歧等特征。"
            "其核心思想是将统计特征与人工表达特征拼接后共同输入分类器，从而提高模型对特定目标场景的解释力。"
            "在实现过程中，本文还引入了对“反对语境”“描述语境”和“直接攻击语境”的区分，以减少仅因出现敏感词而导致的误判。"
            "尽管该模型在总体指标上未超过 BERT，但在人工测试集中的模因式攻击样例上表现出更高的贴题性。"
        ),
        (
            "3.4 BERT 模型实现",
            "BERT 模型部分选用中文预训练模型 bert-base-chinese，并在本任务数据集上进行二分类微调。与传统模型相比，"
            "BERT 不再依赖浅层词频统计，而是通过上下文语义建模学习文本整体含义。本文为 BERT 模型单独设计了配置模块、数据划分模块、训练脚本和批量预测脚本，"
            "并在远程 GPU 环境中完成了正式训练。训练完成后，模型以 Hugging Face 风格目录进行保存，同时生成 metrics.json 和 metadata.json 等结果文件。"
            "实验结果显示，BERT 在正式测试集上的 accuracy、precision、recall 和 F1 指标均优于前两种模型，说明其在总体语义分类能力上具有明显优势。"
        ),
        (
            "3.5 前后端关键实现",
            "在系统实现层面，前端基于 Vue 3 构建，主要负责文本输入、结果展示、状态提示和历史记录显示；后端基于 FastAPI 构建，"
            "主要负责接口注册、模型加载和预测结果返回。前后端之间通过 HTTP 接口进行通信，预测接口以 JSON 形式接收文本并返回分类结果。"
            "对于传统模型，后端通过加载本地 joblib 文件完成预测；对于 BERT 模型，后端和脚本层使用训练后保存的模型目录进行批量预测与对比分析。"
            "为了适应毕业设计场景，系统还实现了健康检查接口、默认模型加载、模型对比脚本以及实验记录文件管理，使系统不仅能够运行，也能够支撑论文中对实验流程和结果的说明。"
        ),
        (
            "3.6 本章小结",
            "本章从数据处理、基线模型、模因式表达特征模型、BERT 模型和前后端实现五个方面介绍了系统的关键实现过程。"
            "通过这一章节可以看出，本文的工作并非仅停留在单一模型训练，而是围绕系统实现、模型演进与实验支撑逐步展开。"
            "第三章的实现内容为后续第四章的实验对比与结果分析提供了直接依据。"
        ),
    ]


def main() -> None:
    draft = find_draft()
    doc = Document(str(draft))

    chapter3_idx = None
    ref_para = None
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text.startswith("第3章"):
            chapter3_idx = i
        if text == "参考文献":
            ref_para = p
            break

    if chapter3_idx is None or ref_para is None:
        raise RuntimeError("未找到第3章占位标题或参考文献位置")

    title_para = doc.paragraphs[chapter3_idx]
    title_para.text = "第3章 关键模型与系统实现"
    format_paragraph(
        title_para,
        font_name="SimSun",
        size_pt=16,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    if chapter3_idx + 1 < len(doc.paragraphs) and not doc.paragraphs[chapter3_idx + 1].text.strip():
        blank = doc.paragraphs[chapter3_idx + 1]._element
        blank.getparent().remove(blank)

    for heading, body in reversed(build_chapter3_items()):
        body_para = ref_para.insert_paragraph_before(body)
        format_paragraph(
            body_para,
            font_name="SimSun",
            size_pt=12,
            bold=False,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_indent_pt=24,
        )
        heading_para = ref_para.insert_paragraph_before(heading)
        format_paragraph(
            heading_para,
            font_name="SimSun",
            size_pt=14,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.LEFT,
        )

    try:
        doc.save(str(draft))
        print(draft)
    except PermissionError:
        fallback = RES_DIR / "draft_with_chapter3.docx"
        doc.save(str(fallback))
        print(fallback)


if __name__ == "__main__":
    main()
