from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_PATH = next((ROOT / "template").glob("*参考模板*.docx"))
OUTPUT_PATH = ROOT / "res" / "初稿.docx"


def set_run_font(run, font_name="宋体", size=12, bold=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(
    doc,
    text="",
    *,
    font_name="宋体",
    size=12,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    first_line_chars=None,
    line_spacing=1.5,
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = line_spacing
    if first_line_chars is not None:
        p.paragraph_format.first_line_indent = Pt(size * first_line_chars)
    run = p.add_run(text)
    set_run_font(run, font_name=font_name, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    if level == 1:
        add_paragraph(
            doc,
            text,
            font_name="宋体",
            size=16,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    else:
        add_paragraph(doc, text, font_name="宋体", size=14, bold=True)


def clear_body_keep_sections(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_table_4_1(doc):
    add_paragraph(
        doc,
        "表4-1 三种模型在正式测试集上的指标对比",
        font_name="黑体",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    table = doc.add_table(rows=4, cols=5)
    headers = ["模型", "Accuracy", "Precision", "Recall", "F1"]
    values = [
        ["Baseline", "0.8223", "0.7754", "0.8025", "0.7887"],
        ["Meme Features", "0.8200", "0.7831", "0.7810", "0.7820"],
        ["BERT", "0.8887", "0.8515", "0.8851", "0.8680"],
    ]
    for i, header in enumerate(headers):
        table.cell(0, i).text = header
    for r, row in enumerate(values, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, size=10)


def build_document():
    doc = Document(str(TEMPLATE_PATH))
    clear_body_keep_sections(doc)

    add_paragraph(
        doc,
        "基于机器学习的文本模因识别系统的研究与设计",
        font_name="宋体",
        size=18,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_paragraph(doc, "摘要", font_name="黑体", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(
        doc,
        "随着社交媒体与网络社区的快速发展，文本内容已经成为网络传播和观点表达的重要载体。"
        "文本模因通常具有模板化、可复制和传播快等特点，其中部分模因式表达会借助群体标签化、"
        "反讽和隐含攻击传播偏见，给网络治理带来新的挑战。本文围绕“基于机器学习的文本模因识别系统的研究与设计”这一课题，"
        "设计并实现了一个面向中文文本的识别系统。系统以前端交互页面和后端预测接口为基础，构建了基线模型、"
        "模因式表达特征增强模型以及BERT模型三条实验路线。实验结果表明，BERT模型在正式测试集上的整体指标最优，"
        "但在人工构造的模因式表达测试集中，模因特征模型对模板化攻击和隐含群体攻击样例更具针对性。"
        "研究说明总体分类性能与特定场景识别能力并不完全一致，在文本模因识别任务中，应结合统计特征、规则特征与语义模型进行综合分析。",
        first_line_chars=2,
    )
    add_paragraph(doc, "关键词：", bold=True)
    add_paragraph(doc, "文本模因识别，机器学习，有害文本检测，BERT，网络内容治理")

    add_paragraph(doc, "")
    add_paragraph(
        doc,
        "Research and Design of a Machine-Learning-Based Text Meme Recognition System",
        font_name="Times New Roman",
        size=16,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    add_paragraph(doc, "Author: Xiao Yang", font_name="Times New Roman", size=12, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_paragraph(doc, "Supervisor: Zhu Jianqi", font_name="Times New Roman", size=12, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_paragraph(doc, "Abstract", font_name="Times New Roman", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(
        doc,
        "This thesis designs a Chinese text recognition system for the task of text meme recognition. "
        "The system includes a front-end interaction page, a back-end prediction service, and three model routes: "
        "a baseline model, a meme-expression-feature-enhanced model, and a BERT-based model. "
        "Experimental results show that BERT achieves the best overall performance on the formal test split, "
        "while the meme-feature model shows stronger sensitivity on manually constructed meme-style attack samples.",
        font_name="Times New Roman",
        size=12,
    )
    add_paragraph(doc, "Keywords:", font_name="Times New Roman", size=12, bold=True)
    add_paragraph(
        doc,
        "text meme recognition, machine learning, harmful text detection, BERT, online content governance",
        font_name="Times New Roman",
        size=12,
    )

    add_paragraph(doc, "")
    add_paragraph(doc, "目    录", font_name="黑体", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for line in [
        "第1章 绪论 .......................................................... 1",
        "第2章 需求分析与总体设计 ............................................ 5",
        "第3章 关键模型与系统实现 ............................................ 10",
        "第4章 实验与结果分析 ................................................ 17",
        "第5章 总结与展望 .................................................... 25",
        "参考文献 ............................................................ 28",
        "致谢 ................................................................. 30",
        "注：提交前请在 Word 中更新目录与页码。",
    ]:
        add_paragraph(doc, line)

    content = {
        "第1章 绪论": {
            "1.1 研究背景": "文本模因是网络空间中具有重复性、可模仿性和传播性的表达形式。随着短文本平台、评论区文化和网络梗语的快速发展，文本模因逐渐成为用户表达态度和传播情绪的重要方式。",
            "1.2 研究意义": "文本模因中的有害表达容易借助模板化传播形成扩散效应。构建一个可运行、可演示并可复现实验的识别系统，不仅具有工程实践价值，也能为后续模因识别研究提供基础。",
            "1.3 国内外研究现状": "现有研究主要集中于有害文本检测、多模态模因检测和仇恨言论识别等方向。总体而言，针对中文文本模因表达的专项研究仍然不足。",
            "1.4 研究内容与论文结构": "本文完成了前后端系统搭建、基线模型实现、模因特征增强模型设计、BERT模型训练与三模型实验比较，并按照系统设计、模型实现和实验分析三条主线展开论述。",
        },
        "第2章 需求分析与总体设计": {
            "2.1 系统目标": "系统需要实现中文文本输入、结果展示、历史记录和模型对比等功能，并支撑实验复现与论文写作。",
            "2.2 功能需求分析": "系统主要包括文本输入模块、文本检测模块、结果展示模块和历史记录模块。用户输入文本后，系统返回分类标签、分数和解释信息。",
            "2.3 非功能需求分析": "系统还需满足可用性、可维护性、可扩展性和实验可复现性要求，便于答辩演示和论文分析。",
            "2.4 总体架构设计": "系统采用前后端分离架构，前端基于 Vue 3，后端基于 FastAPI，模型层包括传统机器学习模型与 BERT 模型，数据层包括训练数据、人工测试集和实验结果文件。",
        },
        "第3章 关键模型与系统实现": {
            "3.1 数据集与预处理": "本文使用包含文本与标签两列的数据集，其中标签0表示正常文本，标签1表示疑似有害文本。训练前完成空值过滤、非法标签处理和重复文本去除，并进行训练集、验证集与测试集划分。",
            "3.2 基线模型设计": "基线模型采用字符级 TF-IDF 与 Logistic Regression 组合，训练速度快、实现简单，适合作为初始对照组。",
            "3.3 模因式表达特征增强模型设计": "为了增强模型对文本模因场景的适配能力，本文在基线模型基础上加入模板化表达、群体指向、情绪强化、嘲讽短语以及语境消歧等特征。",
            "3.4 BERT 模型设计": "BERT 模型选用中文预训练模型 bert-base-chinese，通过二分类微调提升上下文语义理解能力。相比传统模型，BERT 在正式测试集上的总体指标更优。",
            "3.5 系统实现": "系统实现包括前端交互页面、后端接口、传统模型加载、BERT 模型训练与批量预测脚本，并配套实验记录文档和人工测试集。",
        },
        "第4章 实验与结果分析": {
            "4.1 实验环境与评价指标": "传统模型实验在本地开发环境下完成，BERT 正式训练在远程 GPU 环境中完成。评价指标包括准确率、精确率、召回率和 F1 值。",
            "4.2 正式测试集结果": "在正式测试集上，基线模型 F1 为 0.7887，模因特征模型 F1 为 0.7820，BERT 模型 F1 为 0.8680，说明 BERT 在总体语义分类任务上优势明显。",
            "4.3 人工测试集结果": "在人工构造测试集中，BERT 在模因式攻击文本和隐含群体攻击文本上的表现并不占优，而模因特征模型在目标场景上更具针对性。",
            "4.4 结果分析": "实验表明，总体性能与特定场景识别能力并不完全一致。BERT 更适合作为语义分类强基线，而模因特征模型更贴合“文本模因识别”这一课题目标。",
            "4.5 本章小结": "本章通过正式测试集与人工测试集两类实验，对三种模型进行了系统比较，为后续结论和展望提供了依据。",
        },
        "第5章 总结与展望": {
            "5.1 研究总结": "本文完成了文本模因识别系统的设计与实现，构建了三条模型路线，并通过正式测试集和人工测试集进行了比较分析。",
            "5.2 不足与展望": "当前任务仍主要聚焦文本层面的风险表达识别，与完整意义上的多模态模因识别相比尚有差距。未来可进一步引入更贴近网络语境的数据、更适合中文表达的预训练模型，以及图像、上下文会话和传播路径等信息。",
        },
    }

    for chapter, sections in content.items():
        add_heading(doc, chapter, 1)
        for title, text in sections.items():
            add_heading(doc, title, 2)
            add_paragraph(doc, text, first_line_chars=2)
        if chapter == "第4章 实验与结果分析":
            add_table_4_1(doc)

    add_heading(doc, "参考文献", 1)
    references = [
        "[1] Devlin J, Chang M W, Lee K, et al. BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding[C]. Proceedings of NAACL-HLT, 2019.",
        "[2] Baider F, Kozak S A, Strani K, et al. Automatic hate speech detection: a case study on online comments with a focus on self-victimisation and sentiment[J]. Applied Linguistics Review, 2026, 17(2): 641-666.",
        "[3] 陈雅宁, 柯亮, 王文贤, 等. 基于知识增强多任务学习的隐式有害文本检测技术研究[J]. 信息安全研究, 2025, 11(08): 718-726.",
        "[4] 张新生, 张颢泷, 马玉龙, 等. 基于 RoBERTa-MTL 融合语言特征的有害文本识别[J]. 情报杂志, 2026, 45(01): 75-82.",
        "[5] 朱昊. 面向大语言模型的有害文本检测技术的研究与实现[D]. 北京邮电大学, 2024.",
    ]
    for ref in references:
        add_paragraph(doc, ref, size=10)

    add_heading(doc, "致    谢", 1)
    add_paragraph(
        doc,
        "在本次毕业设计的选题、系统实现和论文撰写过程中，指导教师在研究思路、技术路线和写作规范方面给予了耐心指导，使我能够逐步完成从需求分析到实验总结的全过程。同时，也感谢课程学习和项目实践中提供帮助的老师与同学。通过本次毕业设计，我对机器学习文本分类、系统开发流程以及实验分析方法有了更加系统的认识。谨向所有给予帮助和支持的人表示衷心感谢。",
        first_line_chars=2,
    )

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_PATH))
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
